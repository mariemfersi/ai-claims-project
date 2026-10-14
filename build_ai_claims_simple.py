from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

OUT = r"D:\5ds\deep\output\documents\Projet_AI_Claims_Agent_Simplifie.docx"


def font(run, size=10.5, bold=False, color="000000"):
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    font(p.add_run(text))
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    font(p.add_run(text))
    return p


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    tcpr.append(node)


def margins(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side in ("top", "start", "bottom", "end"):
        x = OxmlElement(f"w:{side}")
        x.set(qn("w:w"), "100")
        x.set(qn("w:type"), "dxa")
        mar.append(x)
    tcpr.append(mar)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), "D9D9D9")
        borders.append(e)
    t._tbl.tblPr.append(borders)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.width = Inches(widths[i])
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(c, "1F4E78")
        margins(c)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        font(p.add_run(h), 9.2, True, "FFFFFF")
    trpr = t.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    trpr.append(repeat)
    for rindex, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            c = cells[i]
            c.width = Inches(widths[i])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(c)
            if rindex % 2:
                shade(c, "F3F7FA")
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            font(p.add_run(str(value)), 9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7)
sec.bottom_margin = Inches(0.7)
sec.left_margin = Inches(0.8)
sec.right_margin = Inches(0.8)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")

title = doc.styles["Title"]
title.font.name = "Aptos Display"
title.font.size = Pt(28)
title.font.bold = True
title.font.color.rgb = RGBColor(0, 0, 0)
title.paragraph_format.space_after = Pt(7)

for name, size in (("Heading 1", 17), ("Heading 2", 12.5)):
    s = doc.styles[name]
    s.font.name = "Aptos Display"
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_before = Pt(9)
    s.paragraph_format.space_after = Pt(4)
    s.paragraph_format.keep_with_next = True

p = doc.add_paragraph(style="Title")
p.add_run("AI Claims Agent")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(14)
font(p.add_run("Automatisation intelligente des sinistres automobiles"), 15, False, "1F4E78")

para(doc, "Le projet consiste à créer un workflow n8n qui analyse automatiquement un dossier de sinistre automobile. Le système lit la déclaration et les pièces jointes, vérifie le contrat, recherche les incohérences et prépare une recommandation pour un gestionnaire humain.")

doc.add_heading("Entrées du système", level=1)
for text in [
    "Déclaration écrite du client",
    "Formulaire de sinistre",
    "Photos du véhicule",
    "Constat et facture au format PDF",
    "Déclaration audio en option",
]:
    bullet(doc, text)

doc.add_heading("Architecture retenue", level=1)
para(doc, "L'architecture reste volontairement simple. Un orchestrateur contrôle le workflow. Trois agents travaillent en parallèle, puis deux agents utilisent leurs résultats.")

table(doc, ["Composant", "Rôle"], [
    ["Orchestrateur n8n", "Reçoit le dossier, lance les agents et rassemble leurs sorties"],
    ["Document Agent", "Extrait les informations du formulaire, du constat et de la facture"],
    ["Vision Agent", "Analyse les photos et décrit les dommages visibles"],
    ["Policy RAG Agent", "Vérifie les garanties, exclusions, franchises et délais"],
    ["Fraud Agent", "Compare les résultats et signale les incohérences"],
    ["Decision Agent", "Recommande ACCEPT, INVESTIGATE ou REJECT"],
    ["Report step", "Prépare une synthèse pour le gestionnaire"],
], [1.75, 5.25])

doc.add_heading("Pourquoi cette architecture", level=2)
bullet(doc, "Les agents Document, Vision et Policy peuvent travailler en parallèle, ce qui réduit la latence.")
bullet(doc, "Le Fraud Agent intervient ensuite, car il doit comparer leurs résultats.")
bullet(doc, "Intake et Report restent de simples étapes et non des agents autonomes.")
bullet(doc, "Chaque rôle est facile à expliquer pendant la soutenance.")

doc.add_page_break()

doc.add_heading("Workflow", level=1)
steps = [
    "Le client dépose sa déclaration et ses pièces jointes.",
    "n8n crée le dossier et vérifie les fichiers.",
    "Document Agent, Vision Agent et Policy RAG Agent s'exécutent en parallèle.",
    "Fraud Agent recherche les contradictions entre les documents, les photos et le récit.",
    "Decision Agent produit une recommandation justifiée.",
    "Le gestionnaire consulte le rapport et prend la décision finale.",
]
for item in steps:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    font(p.add_run(item))

doc.add_heading("Recommandations possibles", level=1)
table(doc, ["Résultat", "Signification"], [
    ["ACCEPT", "Le dossier semble couvert et les pièces sont cohérentes"],
    ["INVESTIGATE", "Une information manque ou une incohérence doit être vérifiée"],
    ["REJECT", "Une exclusion claire ou une absence de garantie est retrouvée"],
], [1.45, 5.55])
para(doc, "Ces résultats sont des recommandations. La décision finale reste toujours humaine, surtout pour INVESTIGATE et REJECT.")

doc.add_heading("RAG", level=1)
para(doc, "Le Policy RAG Agent utilise une base documentaire cohérente contenant les conditions générales, les garanties, les exclusions, les franchises, les procédures de gestion et les règles de déclaration.")

doc.add_heading("Comparaison demandée", level=2)
table(doc, ["Méthode", "Configuration"], [
    ["Découpage fixe", "500 tokens avec un overlap de 50 tokens"],
    ["Découpage structurel", "Un chunk par article, garantie ou exclusion"],
], [2.05, 4.95])
para(doc, "Les deux méthodes seront testées sur 15 à 20 questions. Les mesures principales seront Hit at k, MRR et la précision des clauses citées. Chaque réponse devra citer le document et l'article utilisés.")

doc.add_heading("Multimodalité", level=1)
bullet(doc, "PDF : extraction du constat et de la facture.")
bullet(doc, "Images : identification des zones endommagées.")
bullet(doc, "Audio optionnel : transcription de la déclaration.")
para(doc, "La démonstration peut présenter une contradiction multimodale : la déclaration parle d'un choc frontal, mais les photos montrent surtout un dommage arrière.")

doc.add_page_break()

doc.add_heading("Bibliothèque de prompts", level=1)
para(doc, "Chaque prompt aura plusieurs versions. Pour chaque version, l'équipe indiquera son rôle, son format de sortie, l'erreur observée et la correction apportée.")
table(doc, ["Prompt", "Point principal"], [
    ["Document", "Sortie JSON stricte et valeurs absentes indiquées par null"],
    ["Vision", "Description des dommages avec niveau de confiance"],
    ["Policy RAG", "Réponse uniquement à partir des clauses retrouvées"],
    ["Fraud", "Incohérences expliquées avec preuves"],
    ["Decision", "Décision prudente et validation humaine"],
], [1.65, 5.35])

doc.add_heading("Données de test", level=1)
para(doc, "L'équipe créera des dossiers synthétiques afin de ne pas utiliser de vraies données personnelles. Chaque dossier aura une réponse attendue.")
table(doc, ["Cas", "Résultat attendu"], [
    ["Dossier complet et cohérent", "ACCEPT"],
    ["Exclusion clairement applicable", "REJECT"],
    ["Dates contradictoires", "INVESTIGATE"],
    ["Photos incompatibles avec le récit", "INVESTIGATE"],
    ["Facture inhabituelle", "INVESTIGATE"],
    ["Document manquant ou illisible", "Demande de nouvelle pièce"],
], [3.7, 3.3])

doc.add_heading("Évaluation", level=1)
bullet(doc, "Extraction : exactitude des champs et taux de JSON valide.")
bullet(doc, "RAG : Hit at k, MRR et précision des citations.")
bullet(doc, "Fraud : précision des incohérences et nombre de faux positifs.")
bullet(doc, "Décision : matrice de confusion entre résultat attendu et résultat produit.")
bullet(doc, "Workflow : taux de réussite, latence et coût moyen par dossier.")

doc.add_heading("Règles de sécurité", level=1)
bullet(doc, "Utiliser uniquement des données synthétiques pendant le projet.")
bullet(doc, "Ne jamais stocker les clés API dans le dépôt.")
bullet(doc, "Ne pas inventer les informations manquantes.")
bullet(doc, "Exiger des preuves avant une recommandation de rejet.")
bullet(doc, "Garder une validation humaine avant toute décision réelle.")

doc.add_page_break()

doc.add_heading("Planning", level=1)
table(doc, ["Semaine", "Travail attendu"], [
    ["1", "Périmètre, architecture, n8n et dossiers synthétiques"],
    ["2", "Corpus, deux chunkings et premiers tests RAG"],
    ["3", "Agents, prompts v1, vision et PDF en isolation"],
    ["4", "Intégration complète, gestion des erreurs et prompts v2"],
    ["5", "Évaluation, rapport et répétition de la démonstration"],
], [1.0, 6.0])

doc.add_heading("Démonstration", level=1)
for text in [
    "Charger un dossier avec formulaire, constat, facture et photos.",
    "Montrer l'exécution parallèle des agents.",
    "Afficher les informations extraites et les clauses RAG.",
    "Montrer une incohérence détectée.",
    "Présenter la recommandation et le rapport final.",
]:
    bullet(doc, text)

doc.add_heading("Livrables", level=1)
bullet(doc, "Workflow n8n exporté ou dépôt avec instructions d'exécution.")
bullet(doc, "Rapport technique de 8 pages maximum.")
bullet(doc, "Bibliothèque de prompts séparée et versionnée.")
bullet(doc, "Évaluation du retrieval avec questions, résultats et commentaires.")
bullet(doc, "Démonstration de 15 minutes suivie de 10 minutes de questions.")

doc.add_heading("Points à défendre", level=1)
bullet(doc, "Pourquoi utiliser plusieurs agents plutôt qu'un seul appel LLM ?")
bullet(doc, "Pourquoi exécuter trois agents en parallèle ?")
bullet(doc, "Pourquoi le découpage par articles est-il adapté aux contrats ?")
bullet(doc, "Comment le système réduit-il les hallucinations et les faux positifs ?")
bullet(doc, "Pourquoi la décision finale doit-elle rester humaine ?")

# Simple footer
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(footer.add_run("AI Claims Agent - Version simplifiée"), 8.5, False, "666666")

doc.core_properties.title = "AI Claims Agent"
doc.core_properties.subject = "Version simplifiée du projet Advanced Deep Learning"
doc.core_properties.author = "Équipe projet"

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print(OUT)
