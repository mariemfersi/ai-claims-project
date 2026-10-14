from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUTPUT = r"D:\5ds\deep\output\documents\Projet_AI_Claims_Agent.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, name="Aptos", size=10.2, bold=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    set_font(r)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    set_font(r)
    return p


def add_paragraph(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, "1F4E78")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_font(r, size=9.2, bold=True, color="FFFFFF")
        if widths:
            cell.width = Inches(widths[i])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cell = cells[i]
            if ridx % 2 == 1:
                set_cell_shading(cell, "F3F7FA")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(str(text))
            set_font(r, size=9.0)
            if widths:
                cell.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.2)

title = styles["Title"]
title.font.name = "Aptos Display"
title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
title.font.size = Pt(29)
title.font.bold = True
title.font.color.rgb = RGBColor(0, 0, 0)
title.paragraph_format.space_after = Pt(8)

for style_name, size, before, after in (("Heading 1", 17, 11, 5), ("Heading 2", 12.5, 8, 3)):
    s = styles[style_name]
    s.font.name = "Aptos Display"
    s._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    s._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.add_run("AI Claims Agent")
p2 = doc.add_paragraph()
p2.paragraph_format.space_after = Pt(18)
r = p2.add_run("Automatisation intelligente du traitement des sinistres automobiles")
set_font(r, name="Aptos Display", size=16, color="1F4E78")

add_paragraph(doc, "Projet final du module Advanced Deep Learning. Le système reçoit une déclaration de sinistre et ses pièces jointes, extrait les informations utiles, analyse les dommages, vérifie les garanties et produit une recommandation argumentée pour un gestionnaire humain.")

doc.add_heading("Décision de conception", level=1)
add_paragraph(doc, "Le projet adopte une architecture hybride avec un orchestrateur, trois agents spécialisés exécutés en parallèle, puis un agent de cohérence et un agent de décision. Les étapes d'intake et de génération du rapport restent des nœuds du workflow afin de limiter la complexité et de mieux justifier le rôle de chaque agent.")

add_table(doc, ["Élément", "Choix retenu"], [
    ["Plateforme", "n8n comme orchestrateur du workflow"],
    ["Architecture", "Orchestrateur, workers parallèles, contrôle séquentiel"],
    ["Modalités", "Texte, PDF avec mise en page, images, audio en option"],
    ["Décision", "ACCEPT, INVESTIGATE ou REJECT comme recommandation"],
    ["Supervision", "Validation humaine obligatoire avant toute action finale"],
], widths=[1.65, 5.45])

doc.add_heading("Objectifs du projet", level=1)
add_bullet(doc, "Automatiser la réception et la normalisation d'un dossier de sinistre.")
add_bullet(doc, "Extraire des données structurées depuis un formulaire, un constat PDF et une facture.")
add_bullet(doc, "Identifier les zones endommagées et leur sévérité probable à partir des photos.")
add_bullet(doc, "Vérifier les garanties, exclusions, franchises et délais à partir d'une base documentaire RAG.")
add_bullet(doc, "Détecter des incohérences explicables entre la déclaration et les pièces fournies.")
add_bullet(doc, "Produire une recommandation traçable et un rapport synthétique pour le gestionnaire.")

doc.add_page_break()

doc.add_heading("Périmètre fonctionnel", level=1)
doc.add_heading("Entrées du workflow", level=2)
add_table(doc, ["Entrée", "Traitement attendu"], [
    ["Déclaration textuelle", "Identification de l'événement, de la date, du lieu, du véhicule et des parties concernées"],
    ["Formulaire", "Extraction des champs et contrôle des informations obligatoires"],
    ["Constat PDF", "OCR ou lecture native, compréhension de la mise en page et extraction structurée"],
    ["Facture", "Montant, date, fournisseur, lignes principales et cohérence temporelle"],
    ["Photos", "Type de dommage, zone touchée, sévérité estimée et qualité de l'image"],
    ["Audio optionnel", "Transcription, résumé et rapprochement avec la déclaration écrite"],
], widths=[1.55, 5.55])

doc.add_heading("Sorties du workflow", level=2)
add_bullet(doc, "Un objet JSON validé contenant toutes les informations extraites.")
add_bullet(doc, "Les clauses de contrat pertinentes, avec document source et numéro d'article.")
add_bullet(doc, "La liste des incohérences, chacune accompagnée de ses preuves.")
add_bullet(doc, "Une recommandation ACCEPT, INVESTIGATE ou REJECT avec un score de confiance.")
add_bullet(doc, "Un rapport lisible par un gestionnaire et la liste des vérifications humaines nécessaires.")

doc.add_heading("Hors périmètre", level=2)
add_paragraph(doc, "Le prototype ne prononce pas une décision juridique définitive, ne paie pas automatiquement une indemnité, n'identifie pas une personne à partir d'une image et ne prétend pas prouver une fraude. Il signale des incohérences et assiste un gestionnaire humain.")

doc.add_heading("Architecture multi agent", level=1)
add_paragraph(doc, "Les traitements Document, Vision et Policy sont indépendants une fois le dossier créé. Ils sont donc exécutés en parallèle. Le Fraud Agent intervient ensuite, car il compare leurs sorties. Le Decision Agent utilise enfin les faits extraits, les preuves RAG et les incohérences pour établir une recommandation.")

add_table(doc, ["Composant", "Responsabilité", "Outils ou données"], [
    ["Orchestrateur", "Crée le dossier, lance les branches, gère les erreurs et agrège les résultats", "n8n, stockage, validation JSON"],
    ["Document Agent", "Extrait les champs des formulaires, constats et factures", "OCR, parser PDF, schéma JSON"],
    ["Vision Agent", "Décrit les dommages visibles et leur cohérence avec le récit", "Modèle vision, règles de qualité"],
    ["Policy RAG Agent", "Recherche les garanties, exclusions, franchises et délais", "Vector store, métadonnées, reranking"],
    ["Fraud Agent", "Repère les contradictions et calcule un score d'incohérence", "Sorties des autres agents, règles"],
    ["Decision Agent", "Produit une recommandation expliquée et calibrée", "Faits, clauses citées, score d'incohérence"],
    ["Report step", "Génère la synthèse finale sans modifier la décision", "Template de rapport"],
], widths=[1.25, 3.35, 2.5])

doc.add_page_break()

doc.add_heading("Flux de traitement", level=1)
for item in [
    "Le client soumet la déclaration et les pièces jointes depuis un formulaire ou une interface de démonstration.",
    "L'orchestrateur attribue un identifiant, vérifie les fichiers et normalise leurs métadonnées.",
    "Les agents Document, Vision et Policy RAG travaillent en parallèle.",
    "L'orchestrateur valide les sorties avec des schémas JSON et relance une branche en cas d'erreur récupérable.",
    "Le Fraud Agent compare les dates, montants, descriptions, dommages et clauses retrouvées.",
    "Le Decision Agent applique les règles de décision et produit une recommandation avec preuves.",
    "Le Report step prépare le rapport et demande une validation humaine.",
]:
    add_number(doc, item)

doc.add_heading("Règles de décision initiales", level=2)
add_table(doc, ["Recommandation", "Conditions indicatives"], [
    ["ACCEPT", "Garantie retrouvée, aucune exclusion applicable, pièces suffisantes et absence d'incohérence majeure"],
    ["INVESTIGATE", "Information manquante, confiance faible, incohérence significative ou besoin d'expertise humaine"],
    ["REJECT", "Exclusion explicite ou garantie absente, uniquement si la preuve documentaire est claire"],
], widths=[1.45, 5.65])
add_paragraph(doc, "Ces règles servent au prototype. Le gestionnaire conserve la décision finale, en particulier pour les cas INVESTIGATE et REJECT.")

doc.add_heading("Schéma de sortie", level=1)
schema_rows = [
    ["claim_id", "Identifiant unique du dossier", "CLM-2026-001"],
    ["incident", "Date, lieu et description normalisée", "2026-10-06, Tunis"],
    ["vehicle", "Immatriculation et zones endommagées", "rear_bumper, trunk"],
    ["coverage", "Statut, clauses et exclusions", "covered, GAR-ACC-02"],
    ["inconsistencies", "Liste d'écarts avec preuves", "Date facture antérieure"],
    ["recommendation", "ACCEPT, INVESTIGATE ou REJECT", "INVESTIGATE"],
    ["confidence", "Score calibré entre 0 et 1", "0.78"],
    ["human_review_required", "Validation humaine obligatoire", "true"],
]
add_table(doc, ["Champ", "Description", "Exemple"], schema_rows, widths=[1.65, 3.55, 1.9])

doc.add_heading("Gestion des erreurs", level=1)
add_bullet(doc, "Rejeter les types de fichiers non autorisés et limiter leur taille.")
add_bullet(doc, "Conserver le document original et la version extraite.")
add_bullet(doc, "Valider chaque sortie avec un schéma strict avant l'étape suivante.")
add_bullet(doc, "Relancer au maximum une fois une extraction invalide, puis transmettre au gestionnaire.")
add_bullet(doc, "Ne jamais inventer une valeur absente. Utiliser null et signaler l'information manquante.")

doc.add_page_break()

doc.add_heading("Base documentaire RAG", level=1)
doc.add_heading("Corpus", level=2)
add_paragraph(doc, "Le corpus doit rester cohérent. Le prototype utilisera un contrat automobile fictif ou un ensemble de documents provenant du même cadre réglementaire. Il comprendra les conditions générales, les garanties, les exclusions, les franchises, les délais de déclaration et les procédures internes de gestion.")

add_table(doc, ["Document", "Contenu", "Métadonnées"], [
    ["Conditions générales", "Définitions, obligations et conditions", "document, version, article"],
    ["Table des garanties", "Risques couverts et plafonds", "garantie, formule, véhicule"],
    ["Exclusions", "Situations non couvertes", "exclusion, article, catégorie"],
    ["Procédures", "Étapes et pièces nécessaires", "type de sinistre, étape"],
    ["Règles de déclaration", "Délais et justificatifs", "délai, événement, article"],
], widths=[1.55, 3.35, 2.2])

doc.add_heading("Stratégies de découpage", level=2)
add_table(doc, ["Stratégie", "Configuration", "Hypothèse"], [
    ["Découpage fixe", "500 tokens avec overlap de 50 tokens", "Simple mais susceptible de couper un article"],
    ["Découpage structurel", "Un bloc par article, garantie ou exclusion", "Meilleure conservation du contexte juridique"],
], widths=[1.65, 2.65, 2.8])

doc.add_heading("Pipeline de recherche", level=2)
add_bullet(doc, "Nettoyer les documents et préserver les titres, numéros d'articles et tableaux.")
add_bullet(doc, "Créer les embeddings et stocker les chunks avec leurs métadonnées.")
add_bullet(doc, "Appliquer un filtre selon le type de garantie ou de sinistre lorsque cette information est fiable.")
add_bullet(doc, "Récupérer les meilleurs passages avec un top k à tester, par exemple 3, 5 et 8.")
add_bullet(doc, "Reranker les passages si le gain mesuré justifie le coût et la latence.")
add_bullet(doc, "Exiger une citation précise et autoriser la réponse information insuffisante.")

doc.add_heading("Évaluation du RAG", level=2)
add_paragraph(doc, "Un jeu de 15 à 20 requêtes associera chaque question à un ou plusieurs articles attendus. La comparaison portera sur les deux stratégies de découpage, avec les mêmes embeddings et les mêmes requêtes.")
add_table(doc, ["Mesure", "Interprétation"], [
    ["Hit at k", "Présence d'un passage pertinent parmi les k résultats"],
    ["MRR", "Rang du premier passage pertinent"],
    ["Précision des citations", "Clause citée correcte et traçable"],
    ["Fidélité de la réponse", "Conclusion conforme aux passages récupérés"],
    ["Abstention", "Capacité à signaler que le corpus ne permet pas de conclure"],
], widths=[2.05, 5.05])

doc.add_page_break()

doc.add_heading("Multimodalité", level=1)
add_table(doc, ["Modalité", "Usage", "Contrôle"], [
    ["PDF", "Extraction du constat, du formulaire et de la facture", "Comparaison champs extraits et document source"],
    ["Images", "Localisation et description des dommages", "Qualité, visibilité, cohérence avec le récit"],
    ["Audio optionnel", "Transcription de la déclaration orale", "Confiance de transcription et comparaison textuelle"],
], widths=[1.25, 3.25, 2.6])
add_paragraph(doc, "La démonstration doit montrer un cas où la multimodalité apporte une information impossible à obtenir avec le texte seul. Exemple : la déclaration mentionne un choc frontal, tandis que les photos montrent principalement un dommage arrière.")

doc.add_heading("Bibliothèque de prompts", level=1)
add_paragraph(doc, "Chaque prompt sera conservé dans un fichier versionné avec son rôle, ses entrées, son schéma de sortie, la technique utilisée, les cas d'échec observés et les corrections apportées. Les justifications structurées doivent présenter les règles et les preuves sans demander ni stocker un raisonnement interne détaillé.")
add_table(doc, ["Prompt", "Technique", "Échec à tester", "Correction possible"], [
    ["Extraction document", "Sortie JSON et schéma strict", "Champ inventé ou format invalide", "Valeurs null, validation et exemple few shot"],
    ["Analyse vision", "Checklist de zones et confiance", "Dommage affirmé sur photo floue", "Seuil de qualité et abstention"],
    ["Policy RAG", "Réponse fondée sur citations", "Clause non retrouvée ou mal citée", "Refus de conclure sans preuve"],
    ["Fraud", "Règles explicables et preuves croisées", "Faux positif sur un détail mineur", "Niveaux de sévérité et seuils"],
    ["Decision", "Few shot sur cas limites", "Rejet malgré une preuve faible", "Validation humaine et règle de prudence"],
], widths=[1.35, 1.65, 2.0, 2.1])

doc.add_heading("Versionnement des prompts", level=2)
add_bullet(doc, "Nom recommandé : prompts/document_agent/v1.md, v2.md et ainsi de suite.")
add_bullet(doc, "Associer chaque modification à un cas de test ayant échoué.")
add_bullet(doc, "Conserver au moins une comparaison avant et après pour chaque prompt clé.")
add_bullet(doc, "Enregistrer les métriques et les exemples de sortie sans données personnelles réelles.")

doc.add_heading("Données synthétiques", level=1)
add_paragraph(doc, "Le projet utilisera des dossiers entièrement synthétiques afin d'éviter l'emploi de données personnelles ou confidentielles. Chaque dossier possédera une vérité terrain pour les champs, la couverture, les incohérences et la recommandation attendue.")
add_table(doc, ["Scénario", "Caractéristique principale", "Résultat attendu"], [
    ["Dossier simple", "Garantie applicable et pièces cohérentes", "ACCEPT"],
    ["Exclusion", "Clause d'exclusion explicite", "REJECT sous validation humaine"],
    ["Date incohérente", "Dates différentes entre formulaire et constat", "INVESTIGATE"],
    ["Photo contradictoire", "Zone de dommage incompatible avec le récit", "INVESTIGATE"],
    ["Facture suspecte", "Date ou montant incohérent", "INVESTIGATE"],
    ["Document manquant", "Constat ou facture absent", "INVESTIGATE"],
    ["Image inexploitable", "Photo floue ou trop sombre", "Demande de nouvelle pièce"],
    ["Hors corpus", "Question non couverte par les documents", "Abstention"],
], widths=[1.45, 3.75, 1.9])

doc.add_page_break()

doc.add_heading("Évaluation du système", level=1)
add_table(doc, ["Composant", "Mesures"], [
    ["Extraction documentaire", "Exactitude des champs, taux de JSON valide, champs manquants ou inventés"],
    ["Vision", "Exactitude de la zone endommagée, sévérité, taux d'abstention sur images faibles"],
    ["RAG", "Hit at k, MRR, précision des citations, fidélité et abstention"],
    ["Fraud", "Précision et rappel des incohérences, faux positifs par scénario"],
    ["Décision", "Exactitude par classe, matrice de confusion, justification correcte"],
    ["Workflow", "Taux de réussite bout en bout, latence, coût moyen, reprise après erreur"],
], widths=[1.75, 5.35])

doc.add_heading("Critères de réussite proposés", level=2)
add_bullet(doc, "Tous les scénarios de démonstration terminent le workflow sans intervention technique.")
add_bullet(doc, "Chaque recommandation cite les preuves documentaires et les incohérences utilisées.")
add_bullet(doc, "Aucun rejet n'est proposé sans clause explicite ou règle vérifiable.")
add_bullet(doc, "Les sorties invalides sont détectées automatiquement par le schéma.")
add_bullet(doc, "Les cas incertains sont orientés vers INVESTIGATE au lieu de produire une réponse forcée.")

doc.add_heading("Sécurité et gouvernance", level=1)
add_bullet(doc, "Utiliser exclusivement des données synthétiques pendant le projet et la soutenance.")
add_bullet(doc, "Restreindre les formats, la taille et le nombre de fichiers téléversés.")
add_bullet(doc, "Masquer les secrets API dans n8n et ne jamais les inclure dans le dépôt.")
add_bullet(doc, "Journaliser les étapes, versions de prompts, modèles, citations et décisions.")
add_bullet(doc, "Traiter les documents récupérés comme des données, jamais comme des instructions exécutables.")
add_bullet(doc, "Limiter la conservation des fichiers et contrôler les droits d'accès au dossier.")
add_bullet(doc, "Prévoir une validation humaine avant toute décision opérationnelle.")

doc.add_heading("Risques du projet", level=1)
add_table(doc, ["Risque", "Impact", "Réduction"], [
    ["Trop d'agents", "Complexité difficile à expliquer", "Conserver uniquement les rôles avec outils ou décisions propres"],
    ["Corpus incohérent", "Réponses contradictoires", "Utiliser un contrat unique et versionné"],
    ["Hallucination", "Clause ou champ inventé", "Citations obligatoires, schémas et abstention"],
    ["Vision incertaine", "Mauvaise description des dommages", "Contrôle qualité et revue humaine"],
    ["Démo fragile", "Échec réseau ou API", "Scénarios testés, cache et plan de secours"],
    ["Coûts API", "Budget dépassé", "Modèles adaptés, limites et mesure du coût"],
], widths=[1.45, 2.2, 3.45])

doc.add_page_break()

doc.add_heading("Planning sur cinq semaines", level=1)
add_table(doc, ["Semaine", "Travaux", "Livrable de contrôle"], [
    ["1", "Périmètre, architecture, outil, schémas de données et premiers dossiers", "Architecture et scénario de démonstration"],
    ["2", "Corpus, ingestion, deux stratégies de chunking et tests initiaux", "Tableau de comparaison RAG"],
    ["3", "Agents Document, Vision et Policy, prompts v1, tests isolés", "Branches fonctionnelles et résultats"],
    ["4", "Intégration, Fraud, Decision, gestion des erreurs et prompts v2", "Workflow complet et cas d'échec"],
    ["5", "Évaluation, rapport, optimisation, répétition et plan de secours", "Version finale démontrable"],
], widths=[0.75, 4.2, 2.15])

doc.add_heading("Répartition possible pour six membres", level=1)
add_table(doc, ["Responsabilité principale", "Travaux"], [
    ["Orchestration", "n8n, schémas, stockage, gestion des erreurs et intégration"],
    ["Documents", "OCR, extraction PDF et formulaire, validation des champs"],
    ["Vision et audio", "Analyse des photos, qualité des images, transcription optionnelle"],
    ["RAG", "Corpus, chunking, embeddings, retrieval et évaluation"],
    ["Fraud et décision", "Règles d'incohérence, recommandations et supervision humaine"],
    ["Tests et documentation", "Prompts versionnés, dossiers synthétiques, rapport et démonstration"],
], widths=[2.0, 5.1])
add_paragraph(doc, "La répartition organise le travail mais ne limite pas la préparation individuelle. Chaque membre doit pouvoir expliquer l'ensemble du système pendant les questions du jury.")

doc.add_heading("Scénario de démonstration", level=1)
for item in [
    "Soumettre un dossier comprenant un formulaire, un constat PDF, une facture et plusieurs photos.",
    "Afficher le déclenchement du workflow et l'exécution parallèle des trois agents spécialisés.",
    "Montrer les données structurées et les passages RAG avec leurs références.",
    "Présenter une contradiction détectée entre la déclaration et une photo ou une date.",
    "Afficher la recommandation, le score de confiance et les preuves associées.",
    "Terminer par le rapport destiné au gestionnaire et la validation humaine.",
]:
    add_number(doc, item)

doc.add_heading("Jeux de démonstration", level=2)
add_bullet(doc, "Cas ACCEPT avec garantie claire et pièces cohérentes.")
add_bullet(doc, "Cas INVESTIGATE avec une contradiction multimodale.")
add_bullet(doc, "Cas REJECT avec exclusion explicite et preuve documentaire.")

doc.add_page_break()

doc.add_heading("Livrables finaux", level=1)
add_table(doc, ["Livrable", "Contenu minimal"], [
    ["Workflow ou dépôt", "Export n8n ou code, configuration, variables requises et instructions d'exécution"],
    ["Rapport technique", "Huit pages maximum couvrant plateforme, prompts, RAG, multi agent et multimodalité"],
    ["Bibliothèque de prompts", "Versions, rôles, techniques, échecs, corrections et exemples avant après"],
    ["Évaluation retrieval", "Questions, passages attendus, résultats, métriques et commentaires"],
    ["Soutenance", "Démonstration de 15 minutes suivie de 10 minutes de questions"],
], widths=[1.75, 5.35])

doc.add_heading("Plan conseillé du rapport", level=1)
add_table(doc, ["Section", "Contenu"], [
    ["1 Introduction", "Cas d'usage, objectifs, périmètre et limites"],
    ["2 Architecture", "Orchestrateur, agents, flux de données et choix rejetés"],
    ["3 Prompt engineering", "Bibliothèque, versions, échecs et améliorations"],
    ["4 RAG", "Corpus, chunking, embeddings, vector store, retrieval et évaluation"],
    ["5 Multimodalité", "PDF, vision, audio optionnel et contribution au cas d'usage"],
    ["6 Évaluation", "Dossiers synthétiques, métriques, résultats et limites"],
    ["7 Conclusion", "Résultat obtenu et améliorations futures réalistes"],
], widths=[1.65, 5.45])

doc.add_heading("Questions à préparer pour la soutenance", level=1)
add_bullet(doc, "Pourquoi ces composants sont-ils des agents et non de simples appels LLM ?")
add_bullet(doc, "Pourquoi utiliser une architecture hybride plutôt qu'une chaîne entièrement séquentielle ?")
add_bullet(doc, "Comment empêcher une décision sans preuve documentaire suffisante ?")
add_bullet(doc, "Pourquoi le découpage structurel convient-il aux contrats d'assurance ?")
add_bullet(doc, "Comment les métriques RAG ont-elles influencé la configuration finale ?")
add_bullet(doc, "Quels faux positifs le Fraud Agent produit-il et comment sont-ils réduits ?")
add_bullet(doc, "Que se passe-t-il lorsqu'un fichier est illisible ou qu'une API échoue ?")
add_bullet(doc, "Comment protéger les données personnelles dans un déploiement réel ?")

doc.add_heading("Checklist avant remise", level=1)
for text in [
    "Le workflow fonctionne de bout en bout depuis un seul déclencheur.",
    "Les secrets et les données personnelles réelles sont absents du dépôt.",
    "Chaque agent possède un rôle justifié, des entrées et une sortie validée.",
    "Les prompts sont versionnés et contiennent des exemples avant et après.",
    "Les deux stratégies de chunking sont comparées sur les mêmes questions.",
    "Les recommandations comportent des preuves et une validation humaine.",
    "Les trois scénarios de démonstration ont été répétés.",
    "Tous les membres peuvent expliquer chaque partie du système.",
]:
    add_bullet(doc, "☐ " + text)

# Header and footer
header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = hp.add_run("Advanced Deep Learning  |  AI Claims Agent")
set_font(hr, size=8.5, color="5A5A5A")

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run("Page ")
set_font(run, size=8.5, color="666666")
fld_char1 = OxmlElement("w:fldChar")
fld_char1.set(qn("w:fldCharType"), "begin")
instr = OxmlElement("w:instrText")
instr.set(qn("xml:space"), "preserve")
instr.text = " PAGE "
fld_char2 = OxmlElement("w:fldChar")
fld_char2.set(qn("w:fldCharType"), "end")
run._r.append(fld_char1)
run._r.append(instr)
run._r.append(fld_char2)

# Metadata
doc.core_properties.title = "AI Claims Agent"
doc.core_properties.subject = "Projet final Advanced Deep Learning"
doc.core_properties.author = "Équipe projet"
doc.core_properties.keywords = "deep learning, agents, RAG, multimodalité, assurance, n8n"

import os
os.makedirs(os.path.dirname(OUTPUT), exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
